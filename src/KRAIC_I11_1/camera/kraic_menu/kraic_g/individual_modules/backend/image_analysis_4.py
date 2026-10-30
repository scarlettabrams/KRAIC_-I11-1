import os
import re
import shutil

import openpyxl

VALID_EXTS = (".jpg", ".jpeg", ".png", ".tiff", ".bmp")


def safe_makedirs(path):
    os.makedirs(path, exist_ok=True)


def move_device_folders(flir_device_path, image_save_folder):
    """
    Move top-level FLIR folders into image_save_folder.
    """
    safe_makedirs(image_save_folder)

    for name in os.listdir(flir_device_path):
        src = os.path.join(flir_device_path, name)

        if not os.path.isdir(src):
            continue

        dst = os.path.join(image_save_folder, name)

        if os.path.exists(dst):
            base = name
            i = 1
            while os.path.exists(os.path.join(image_save_folder, f"{base}_{i}")):
                i += 1
            dst = os.path.join(image_save_folder, f"{base}_{i}")

        shutil.move(src, dst)

    return image_save_folder


def image_sort_key(path):
    """
    Sort FLIR1600, FLIR1601, FLIR1602... by the number in the filename.
    """
    fname = os.path.basename(path)
    m = re.search(r"(\d+)", fname)
    return int(m.group(1)) if m else fname


def gather_and_sort_images(root):
    files = []

    for dirpath, _, filenames in os.walk(root):
        for f in filenames:
            if f.lower().endswith(VALID_EXTS):
                files.append(os.path.join(dirpath, f))

    files.sort(key=image_sort_key)
    return files


def extract_temps_for_images(image_paths, coordinates, orientation):
    try:
        from flirimageextractor import FlirImageExtractor
    except ImportError as e:
        raise ImportError(
            "Could not import flirimageextractor on this Raspberry Pi."
        ) from e

    fle = FlirImageExtractor()
    results = {}

    x = int(coordinates["x"])
    y_top = int(coordinates["y_top"])
    y_bottom = int(coordinates["y_bottom"])

    if orientation == "downward":
        # lowest vertical point first, highest last
        y_values = range(y_bottom, y_top - 1, -1)

    elif orientation == "upward":
        # highest vertical point first, lowest last
        y_values = range(y_top, y_bottom + 1)

    else:
        raise ValueError("orientation must be either 'upward' or 'downward'")

    for image_path in image_paths:
        fname = os.path.basename(image_path)

        fle.process_image(image_path)
        thermal = fle.get_thermal_np()

        temps = []

        for y in y_values:
            if 0 <= y < thermal.shape[0] and 0 <= x < thermal.shape[1]:
                temps.append(float(thermal[y, x]))
            else:
                temps.append(None)

        results[fname] = temps

    return results, len(list(y_values))


def write_temperatures_to_workbook(cal_wb_path, image_temps):
    wb = openpyxl.load_workbook(cal_wb_path)
    ws = wb["Calibration Results"]

    start_col = 8  # column H

    for row, (img_name, temps) in enumerate(image_temps.items(), start=2):
        for i, temp in enumerate(temps):
            ws.cell(row=row, column=start_col + i, value=temp)

        image_name_col = start_col + len(temps) + 3
        ws.cell(row=row, column=image_name_col, value=img_name)

    wb.save(cal_wb_path)


def main(
    flir_device_path,
    image_save_folder,
    coordinates,
    wb_path,
    cal_progress,
    folder_name,
    orientation,
):
    try:
        move_device_folders(flir_device_path, image_save_folder)

        image_paths = gather_and_sort_images(image_save_folder)

        print("\nImages found:")

        for img in image_paths:
            print(os.path.basename(img))

        print("\nImages found:")

        for img in image_paths:
            print(os.path.basename(img))

        if not image_paths:
            raise FileNotFoundError(
                "No calibration images found after moving FLIR folders."
            )

        temps, number = extract_temps_for_images(image_paths, coordinates, orientation)

        write_temperatures_to_workbook(wb_path, temps)

        return "Images processed successfully", True

    except Exception as e:
        print(f"Error during image processing: {e}")
        return f"Error: {e}", False
