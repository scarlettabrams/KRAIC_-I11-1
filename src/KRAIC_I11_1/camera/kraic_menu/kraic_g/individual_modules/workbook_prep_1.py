from openpyxl import load_workbook
from openpyxl.styles import Font

from KRAIC_I11_1.camera.helper.helper import (
    load_json,
    update_json,
)


def prepare_calibration_workbook(folder_name, wb_path, cal_progress):

    print("\nPreparing calibration workbook...\n")

    while True:
        inputs = collect_inputs()

        if confirm_inputs(inputs):
            write_to_excel(inputs, wb_path)
            print("\nCalibration workbook saved")
            break
        else:
            print("\nLet's try again.\n")

    wb = load_workbook(wb_path)
    # avoid duplicate sheet names
    if "Calibration Results" in wb.sheetnames:
        ws = wb["Calibration Results"]
    else:
        ws = wb.create_sheet("Calibration Results")

    progress = load_json(cal_progress)
    state = progress["CALIBRATION_STATUS"][folder_name]

    while True:
        print(
            "Please define the orientation of the KRAIC-G by telling the program if "
            "crystallising solution moves down or up the height of the KRAIC-G "
        )

        orientation = input(
            "\nSelect 'd' if crystallising solution moves downward and 'u' if "
            "crystallising solution moves upwards:"
        ).lower()

        if orientation == "d" or "u":
            break
        else:
            print(f"'{orientation}' is an invalid selection, please try again")
            continue

    if orientation == "u":
        direction = "upward"

    else:
        direction = "downward"

    state["Orientation"] = direction

    coordinate_1 = state["Coordinates"].get("y_top")
    coordinate_2 = state["Coordinates"].get("y_bottom")

    n = abs(coordinate_1 - coordinate_2) + 1

    titles = (
        [
            "Experiment Number",
            "Flowrate (ml/min)",
            "Julabo Setpoint(°C)",
            "Heating PXU Setpoint(°C)",
            "Julabo Temperature(°C)",
            "Coolant Temperature at Inlet(°C)",
            "Cold Extremity Temperature(°C)",
        ]
        + n * [""]
        + [
            "Hot Extremity Temperature(°C)",
            "Coolant Temperature at Outlet(°C)",
            "Ambient Temperature",
            "Image Name",
        ]
    )

    ws.append(titles)
    wb.save(wb_path)

    state["Preparation"] = True

    progress = update_json(cal_progress, progress)

    print("\n\nWorkbook preparation complete - calibrations ready to begin")


def collect_inputs():
    """Prompt user for all fields and return dict."""

    fields = [
        "Location:",
        "Water storage vessel temperature (°C):",
        "Storage vessel to pump ID (inches)",
        "Storage vessel to pump length (m):",
        "Pump to cooling jacket ID(inches):",
        "Pump to cooling jacket length (m):",
        "ID of tubing inside cooling jacket (inches):",
        "Length of tubing inside cooling jacket (m):",
        "ID of tubing from end of cooling Jacket to AM adapter(inches):",
        "Length of tubing from end of cooling jacket to AM adapter(m):",
        "ID of tubing from AM adapter to the KRAIC-G(inches):",
        "Length of tubing from AM adapter to the KRAIC-G(m):",
        "ID of tubing from the KRAIC-G to the glass adapter(inches):",
        "Length of tubing from the KRAIC-G to the glass adapter(m):",
        "ID of tubing from the glass adapter to the storage vessel(inches):",
        "Length of tubing from the glass adapter to the storage vessel(m):",
        "Seg-bath temperature:",
        "Ambient temperature:",
        "Additional Information:",
    ]

    data = {}
    print("\nPlease enter calibration details:\n")

    for field in fields:
        data[field] = input(f"{field} ")

    return data


def confirm_inputs(data):
    """Display inputs and ask for confirmation."""
    print("\nPlease review your entries:\n")

    for field, value in data.items():
        print(f"{field} {value}")

    while True:
        confirm = input("\nAre you happy with these entries? (y/n): ").lower()
        if confirm in ("y", "n"):
            return confirm == "y"
        print("Please enter 'y' or 'n'.")


def write_to_excel(data, wb_path):

    calibration_title = "KRAIC-G Calibration"

    wb = load_workbook(wb_path)
    # avoid duplicate sheet names
    if "Calibration Description" in wb.sheetnames:
        ws = wb["Calibration Description"]
    else:
        ws = wb.create_sheet("Calibration Description")

    bold = Font(bold=True)

    # Write calibration title
    ws["A1"] = "Calibration Title:"
    ws["A1"].font = bold
    ws["B1"] = calibration_title

    row = 2
    for field, value in data.items():
        ws.cell(row=row, column=1, value=field).font = bold
        ws.cell(row=row, column=2, value=value)
        row += 1

    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 30

    wb.save(wb_path)
